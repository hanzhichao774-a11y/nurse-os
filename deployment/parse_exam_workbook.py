#!/usr/bin/env python3
"""Convert the supplied teacher qualification workbook to the demo JSON format."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


HEADER_RE = re.compile(
    r"^(高等教育学|高等教育心理学|高校教师职业道德修养和高等教育法规)\s*"
    r"(单选|多选|判断)\s*(\d+)\s*题$"
)
NUMBER_RE = re.compile(r"^\s*(\d+)\s*[\.．、]\s*")
OPTION_RE = re.compile(r"(?:^|\s)([A-E])[\.．、]\s*")

SUBJECT_NAMES = {
    "高等教育学": "高等教育学",
    "高等教育心理学": "高等教育心理学",
    "高校教师职业道德修养和高等教育法规": "职业道德与法规",
}
TYPE_NAMES = {"单选": "single", "多选": "multiple", "判断": "judge"}


def compact(value: object) -> str:
    text = str(value or "").replace("\u3000", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_choice(raw: str) -> tuple[int, str, list[dict[str, str]]]:
    match = NUMBER_RE.match(raw)
    if not match:
        raise ValueError(f"题目缺少序号: {raw[:80]}")
    number = int(match.group(1))
    body = raw[match.end() :]
    markers = list(OPTION_RE.finditer(body))
    if len(markers) < 2:
        raise ValueError(f"选项不足: {raw[:120]}")
    question = compact(body[: markers[0].start()])
    options: list[dict[str, str]] = []
    for index, marker in enumerate(markers):
        end = markers[index + 1].start() if index + 1 < len(markers) else len(body)
        options.append({"key": marker.group(1), "text": compact(body[marker.end() : end])})
    return number, question, options


def parse_judge(raw: str) -> tuple[int, str]:
    match = NUMBER_RE.match(raw)
    if not match:
        raise ValueError(f"判断题缺少序号: {raw[:80]}")
    return int(match.group(1)), compact(raw[match.end() :])


def build(source: Path) -> tuple[dict, dict]:
    # Normal mode is deliberate: repeated random access on a read-only worksheet
    # reparses the XML stream and turns this small workbook into a minute-long job.
    sheet = load_workbook(source, read_only=False, data_only=True).active
    sections: list[dict] = []
    for row_number, row in enumerate(sheet.iter_rows(min_col=1, max_col=3, values_only=True), 1):
        for value in row:
            if not value:
                continue
            match = HEADER_RE.match(compact(value))
            if match:
                sections.append(
                    {
                        "row": row_number,
                        "subject_raw": match.group(1),
                        "type_raw": match.group(2),
                        "expected": int(match.group(3)),
                    }
                )
                break
    # The source repeats its first section title on two adjacent rows.
    sections = [section for index, section in enumerate(sections) if not index or section != sections[index - 1]]
    deduped: list[dict] = []
    for section in sections:
        if deduped and all(section[k] == deduped[-1][k] for k in ("subject_raw", "type_raw", "expected")):
            deduped[-1]["row"] = section["row"]
        else:
            deduped.append(section)
    sections = deduped

    questions: list[dict] = []
    repairs: list[dict] = []
    section_report: list[dict] = []
    for section_index, section in enumerate(sections):
        start = section["row"] + 1
        end = sections[section_index + 1]["row"] - 1 if section_index + 1 < len(sections) else sheet.max_row
        subject = SUBJECT_NAMES[section["subject_raw"]]
        question_type = TYPE_NAMES[section["type_raw"]]
        parsed_count = 0
        seen_numbers: list[int] = []
        for row_number in range(start, end + 1):
            a = sheet.cell(row_number, 1).value
            b = sheet.cell(row_number, 2).value
            c = sheet.cell(row_number, 3).value
            raw = next((str(value) for value in (b, c) if value not in (None, "")), "")
            if not raw:
                continue
            if question_type == "judge":
                number, prompt = parse_judge(raw)
                answer_value = compact(a)
                if not answer_value and subject == "高等教育学" and number == 42:
                    answer_value = "✔"
                    repairs.append(
                        {
                            "sheetRow": row_number,
                            "question": number,
                            "field": "answer",
                            "original": None,
                            "repaired": "正确",
                            "reason": "原表答案单元格为空；题干与同表关于德育过程的定义完全一致。",
                        }
                    )
                if answer_value not in {"✔", "✘"}:
                    raise ValueError(f"第 {row_number} 行判断题答案无效: {answer_value!r}")
                answer = "T" if answer_value == "✔" else "F"
                options = [{"key": "T", "text": "正确"}, {"key": "F", "text": "错误"}]
            else:
                number, prompt, options = parse_choice(raw)
                answer = compact(a)
                valid_keys = {item["key"] for item in options}
                if not answer or any(key not in valid_keys for key in answer):
                    raise ValueError(f"第 {row_number} 行答案 {answer!r} 不在选项 {sorted(valid_keys)} 中")
                if question_type == "single" and len(answer) != 1:
                    raise ValueError(f"第 {row_number} 行单选题有多个答案: {answer}")
                if question_type == "multiple" and len(answer) < 2:
                    raise ValueError(f"第 {row_number} 行多选题答案不足: {answer}")
            seen_numbers.append(number)
            parsed_count += 1
            questions.append(
                {
                    "id": f"q{len(questions) + 1:04d}",
                    "subject": subject,
                    "type": question_type,
                    "number": number,
                    "prompt": prompt,
                    "options": options,
                    "answer": answer,
                }
            )
        if parsed_count != section["expected"]:
            raise ValueError(
                f"{subject}/{question_type} 应有 {section['expected']} 题，实际解析 {parsed_count} 题"
            )
        expected_numbers = list(range(1, section["expected"] + 1))
        if seen_numbers != expected_numbers:
            raise ValueError(f"{subject}/{question_type} 题号不连续")
        section_report.append(
            {
                "subject": subject,
                "type": question_type,
                "expected": section["expected"],
                "parsed": parsed_count,
                "sourceRows": [start, end],
            }
        )

    counts_by_subject = Counter(question["subject"] for question in questions)
    counts_by_type = Counter(question["type"] for question in questions)
    fingerprint = f"{source.stat().st_size}-{int(source.stat().st_mtime)}"
    payload = {
        "meta": {
            "title": "高校教师资格笔试题库",
            "version": "2026.09",
            "total": len(questions),
            "sourceFingerprint": fingerprint,
            "subjects": dict(counts_by_subject),
            "types": dict(counts_by_type),
        },
        "questions": questions,
    }
    report = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "source": source.name,
        "sheet": sheet.title,
        "result": "passed",
        "totalExpected": 1954,
        "totalParsed": len(questions),
        "sections": section_report,
        "repairs": repairs,
        "checks": {
            "allQuestionsHavePrompt": all(bool(question["prompt"]) for question in questions),
            "allQuestionsHaveAnswer": all(bool(question["answer"]) for question in questions),
            "allChoiceAnswersExistInOptions": all(
                all(key in {option["key"] for option in question["options"]} for key in question["answer"])
                for question in questions
            ),
            "idsAreUnique": len({question["id"] for question in questions}) == len(questions),
        },
    }
    if len(questions) != 1954 or not all(report["checks"].values()):
        raise ValueError("题库完整性校验失败")
    return payload, report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    payload, report = build(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"total": payload["meta"]["total"], "output": str(args.output), "report": str(args.report)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
