#!/usr/bin/env python3
"""Build the political-worker exam question bank from the supplied Word files.

The source files contain the answer key inline:

* fill-in answers are underlined Word runs;
* single-choice answers are the final parenthesized phrase;
* multiple-choice answers are all parenthesized phrases.

Because the source has no distractors, this script creates deterministic distractors
from answers in the same syllabus section. The generated JSON therefore remains
stable across deployments and can be reviewed or regenerated at any time.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
QUESTION_RE = re.compile(r"^(\d+)[.．、]\s*(.+)$", re.S)
SECTION_RE = re.compile(r"^[一二三四五六七八九十]+[、.]\s*(.+?)(?:部分)?$")
PAREN_RE = re.compile(r"[（(]([^()（）]+)[）)]")

SECTION_NAMES = {
    "党的二十大报告": "党的二十大报告",
    "党的二十届四中全会《建议》": "党的二十届四中全会《建议》",
    "《中国共产党章程》": "中国共产党章程",
    "《中华人民共和国宪法》": "中华人民共和国宪法",
    "《中华人民共和国爱国主义教育法》": "中华人民共和国爱国主义教育法",
    "《中国共产党思想政治工作条例》（以下简称《条例》）": "中国共产党思想政治工作条例",
    "《中国共产党思想政治工作条例》": "中国共产党思想政治工作条例",
    "《关于在全党开展树立和践行正确政绩观学习教育的通知》《中国共产党河北省第十届委员会第十次全体会议决议》": "政绩观学习教育与河北省委全会",
}


@dataclass
class RawQuestion:
    source_type: str
    number: int
    subject: str
    prompt: str
    answers: list[str]
    source_text: str


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", "", value).strip()


def clean_answer(value: str) -> str:
    value = normalize_space(value)
    return value.strip("，。；;、：:！？!?（）() \t\r\n")


def is_underlined(run) -> bool:
    underline = run._r.find(f"{{{W_NS}}}rPr/{{{W_NS}}}u")
    if underline is None:
        return False
    value = underline.get(f"{{{W_NS}}}val", "single")
    return value not in {"none", "0", "false"}


def cloze_from_runs(paragraph) -> tuple[str, list[str]]:
    """Replace contiguous underlined runs with a blank and return their text."""
    parts: list[str] = []
    answers: list[str] = []
    pending: list[str] = []

    def flush() -> None:
        if not pending:
            return
        answer = clean_answer("".join(pending))
        if answer:
            answers.append(answer)
            parts.append("（　）")
        pending.clear()

    for run in paragraph.runs:
        if is_underlined(run) and run.text:
            pending.append(run.text)
        else:
            flush()
            parts.append(run.text)
    flush()
    text = normalize_space("".join(parts))
    match = QUESTION_RE.match(text)
    return (match.group(2) if match else text), answers


def normalize_section(text: str) -> str | None:
    match = SECTION_RE.match(normalize_space(text))
    if not match:
        return None
    raw = match.group(1).removesuffix("部分")
    return SECTION_NAMES.get(raw, raw.strip("《》"))


def split_merged_questions(text: str) -> list[str]:
    """Split the one source paragraph where questions 16 and 17 were merged."""
    first = QUESTION_RE.match(text)
    if not first:
        return [text]
    current = int(first.group(1))
    marker = re.search(rf"(?<=[。！？!])\s*({current + 1})[.．、]", text)
    if not marker:
        return [text]
    return [text[: marker.start()].strip(), text[marker.start():].strip()]


def parse_fill(path: Path) -> list[RawQuestion]:
    document = Document(path)
    subject = "未分类"
    result: list[RawQuestion] = []
    for paragraph in document.paragraphs:
        text = normalize_space(paragraph.text)
        if not text:
            continue
        heading = normalize_section(text)
        if heading:
            subject = heading
            continue
        match = QUESTION_RE.match(text)
        if not match:
            continue
        prompt, answers = cloze_from_runs(paragraph)
        if not answers:
            raise ValueError(f"填空题未找到下划线答案：{text}")
        # Formatting can split one answer into several adjacent runs. The source
        # contains one intended blank per numbered question.
        answer = clean_answer("".join(answers))
        result.append(RawQuestion("fill", int(match.group(1)), subject, prompt, [answer], match.group(2)))
    return result


def parse_parenthesized(path: Path, source_type: str) -> list[RawQuestion]:
    document = Document(path)
    subject = "未分类"
    result: list[RawQuestion] = []
    for paragraph in document.paragraphs:
        text = normalize_space(paragraph.text)
        if not text:
            continue
        heading = normalize_section(text)
        if heading:
            subject = heading
            continue
        for piece in split_merged_questions(text):
            match = QUESTION_RE.match(piece)
            if not match:
                continue
            body = match.group(2)
            found = [clean_answer(item) for item in PAREN_RE.findall(body)]
            found = [item for item in found if item]
            if not found:
                raise ValueError(f"{source_type}题未找到括号答案：{piece}")
            answers = [found[-1]] if source_type == "single" else list(dict.fromkeys(found))
            prompt = body
            for answer in answers:
                prompt = re.sub(rf"[（(]{re.escape(answer)}[）)]", "（　）", prompt, count=1)
            result.append(RawQuestion(source_type, int(match.group(1)), subject, prompt, answers, body))
    return result


def stable_rng(question: RawQuestion) -> random.Random:
    seed_text = f"{question.source_type}|{question.subject}|{question.number}|{question.source_text}"
    seed = int(hashlib.sha256(seed_text.encode("utf-8")).hexdigest()[:16], 16)
    return random.Random(seed)


def choose_distractors(question: RawQuestion, pool: list[str], count: int) -> list[str]:
    answers = set(question.answers)
    candidates = [item for item in dict.fromkeys(pool) if item not in answers and item]
    rng = stable_rng(question)
    rng.shuffle(candidates)
    target_len = sum(map(len, question.answers)) / len(question.answers)
    candidates.sort(key=lambda item: (abs(len(item) - target_len), rng.random()))
    return candidates[:count]


def make_options(question: RawQuestion, section_pool: list[str]) -> tuple[list[dict[str, str]], str]:
    if question.source_type in {"single", "fill"}:
        distractors = choose_distractors(question, section_pool, 3)
    else:
        # Keep the option set readable on phones while ensuring at least one
        # distractor exists for every multiple-choice question.
        desired = max(1, min(2, 6 - len(question.answers)))
        distractors = choose_distractors(question, section_pool, desired)
    values = list(dict.fromkeys([*question.answers, *distractors]))
    if len(values) < len(question.answers) + 1:
        raise ValueError(f"无法为题目生成有效干扰项：{question.source_text}")
    rng = stable_rng(question)
    rng.shuffle(values)
    keys = "ABCDEFGH"
    options = [{"key": keys[index], "text": value} for index, value in enumerate(values)]
    answer_keys = "".join(sorted(keys[values.index(answer)] for answer in question.answers))
    return options, answer_keys


def validate(raw: list[RawQuestion], questions: list[dict]) -> None:
    if len(raw) != 660:
        raise ValueError(f"题目总数异常：期望 660，实际 {len(raw)}")
    expected = {"fill": 232, "single": 272, "multiple": 156}
    actual = Counter(item.source_type for item in raw)
    if dict(actual) != expected:
        raise ValueError(f"题型数量异常：{dict(actual)}")
    ids = [item["id"] for item in questions]
    if len(ids) != len(set(ids)):
        raise ValueError("题目 ID 重复")
    for item in questions:
        keys = {option["key"] for option in item["options"]}
        if not item["answer"] or not set(item["answer"]) <= keys:
            raise ValueError(f"答案与选项不匹配：{item['id']}")
        if item["type"] in {"single", "fill"} and len(item["answer"]) != 1:
            raise ValueError(f"单答案题出现多个答案：{item['id']}")
        if item["type"] == "multiple" and len(item["answer"]) < 2:
            raise ValueError(f"多选题答案不足两个：{item['id']}")


def build(fill_path: Path, single_path: Path, multiple_path: Path) -> dict:
    raw = [
        *parse_fill(fill_path),
        *parse_parenthesized(single_path, "single"),
        *parse_parenthesized(multiple_path, "multiple"),
    ]
    pools: dict[str, list[str]] = defaultdict(list)
    for item in raw:
        pools[item.subject].extend(item.answers)

    questions: list[dict] = []
    for index, item in enumerate(raw, 1):
        # Some fill-in source sentences repeat the answer before the underlined
        # blank (for example, the same phrase is used as both a feature and an
        # advantage). Hide every remaining occurrence so the prompt cannot
        # reveal its own answer.
        prompt = item.prompt
        for answer_text in item.answers:
            prompt = prompt.replace(answer_text, "（　）")
        options, answer = make_options(item, pools[item.subject])
        questions.append(
            {
                "id": f"zg{index:04d}",
                "subject": item.subject,
                "type": item.source_type,
                "number": item.number,
                "prompt": prompt,
                "options": options,
                "answer": answer,
                "sourceText": item.source_text,
            }
        )
    validate(raw, questions)
    subjects = Counter(item.subject for item in raw)
    types = Counter(item.source_type for item in raw)
    source_fingerprint = hashlib.sha256(
        b"".join(path.read_bytes() for path in (fill_path, single_path, multiple_path))
    ).hexdigest()[:16]
    return {
        "meta": {
            "title": "政工师专业知识题库",
            "version": "2026.10",
            "total": len(questions),
            "sourceFingerprint": source_fingerprint,
            "subjects": dict(subjects),
            "types": dict(types),
        },
        "questions": questions,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fill", type=Path, required=True)
    parser.add_argument("--single", type=Path, required=True)
    parser.add_argument("--multiple", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    bank = build(args.fill, args.single, args.multiple)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(bank, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(bank["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
